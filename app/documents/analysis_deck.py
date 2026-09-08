from pathlib import Path
from datetime import datetime, timedelta

from pptx import Presentation
from pptx.util import Inches, Pt

from sqlalchemy import select

from app.infrastructure.database.database import SessionLocal
from app.infrastructure.database.models import Bill


def generate_sales_analysis_deck(days: int = 7) -> str:

    if days <= 0:
        raise ValueError("DAYS_MUST_BE_POSITIVE")

    output_dir = Path("data/reports")
    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    db = SessionLocal()

    try:
        end = datetime.utcnow()

        start = end - timedelta(days=days)

        bills = db.scalars(
            select(Bill).where(
                Bill.status == "FINALIZED",
                Bill.finalized_at >= start,
                Bill.finalized_at <= end
            )
        ).all()

        total_sales = sum(
            bill.total
            for bill in bills
        )

        total_bills = len(bills)

        average_bill = (
            total_sales / total_bills
            if total_bills > 0
            else 0
        )

        daily_sales = {}

        for i in range(days):

            current_day = (
                start + timedelta(days=i)
            ).date()

            daily_sales[current_day] = 0

        for bill in bills:

            if not bill.finalized_at:
                continue

            bill_day = bill.finalized_at.date()

            if bill_day in daily_sales:
                daily_sales[bill_day] += bill.total

        presentation = Presentation()

        # --------------------------------------------------
        # TITLE SLIDE
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[0]
        )

        slide.shapes.title.text = (
            "KiranaPilot Sales Analysis"
        )

        slide.placeholders[1].text = (
            f"Sales performance for the last {days} days"
        )

        # --------------------------------------------------
        # SUMMARY SLIDE
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        )

        slide.shapes.title.text = "Sales Summary"

        summary_text = (
            f"Total Sales: ₹{total_sales:.2f}\n\n"
            f"Total Bills: {total_bills}\n\n"
            f"Average Bill Value: ₹{average_bill:.2f}\n\n"
            f"Analysis Period: {days} days"
        )

        textbox = slide.shapes.add_textbox(
            Inches(1),
            Inches(1.7),
            Inches(8),
            Inches(4)
        )

        text_frame = textbox.text_frame

        paragraph = text_frame.paragraphs[0]

        paragraph.text = summary_text
        paragraph.font.size = Pt(24)

        # --------------------------------------------------
        # DAILY SALES SLIDE
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        )

        slide.shapes.title.text = "Daily Sales"

        rows = len(daily_sales) + 1

        table = slide.shapes.add_table(
            rows,
            2,
            Inches(1),
            Inches(1.5),
            Inches(7),
            Inches(4.5)
        ).table

        table.cell(0, 0).text = "Date"
        table.cell(0, 1).text = "Sales"

        for row, (date, sales) in enumerate(
            daily_sales.items(),
            start=1
        ):
            table.cell(row, 0).text = (
                date.strftime("%d-%m-%Y")
            )

            table.cell(row, 1).text = (
                f"₹{sales:.2f}"
            )

        # --------------------------------------------------
        # INSIGHTS SLIDE
        # --------------------------------------------------

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        )

        slide.shapes.title.text = "Key Insights"

        highest_day = None
        highest_sales = 0

        for date, sales in daily_sales.items():

            if sales > highest_sales:
                highest_sales = sales
                highest_day = date

        if highest_day:

            best_day_text = (
                f"Highest sales day: "
                f"{highest_day.strftime('%d-%m-%Y')}\n"
                f"Sales: ₹{highest_sales:.2f}"
            )

        else:

            best_day_text = (
                "No finalized sales were recorded "
                "during this period."
            )

        insight_text = (
            f"{best_day_text}\n\n"
            f"Total finalized bills: {total_bills}\n\n"
            f"Average bill value: ₹{average_bill:.2f}\n\n"
            "Only finalized bills are included "
            "in this analysis."
        )

        textbox = slide.shapes.add_textbox(
            Inches(1),
            Inches(1.5),
            Inches(8),
            Inches(4.5)
        )

        paragraph = (
            textbox.text_frame.paragraphs[0]
        )

        paragraph.text = insight_text
        paragraph.font.size = Pt(20)

        # --------------------------------------------------
        # SAVE
        # --------------------------------------------------

        file_path = (
            output_dir
            / (
                "sales_analysis_"
                f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"
                ".pptx"
            )
        )

        presentation.save(str(file_path))

        return str(file_path)

    finally:
        db.close()